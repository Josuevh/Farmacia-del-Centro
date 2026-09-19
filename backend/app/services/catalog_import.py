import csv
import io
import re
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy.exc import IntegrityError
from sqlalchemy.future import select

from app import models

# Spanish (and a couple English) header aliases we recognize, case/accent-insensitive.
COLUMN_ALIASES = {
    'name': ['nombre', 'name', 'producto'],
    'sku': ['sku'],
    'barcode': ['codigo_barras', 'codigo de barras', 'codigo barras', 'barcode'],
    'price': ['precio', 'price'],
    'category': ['categoria', 'category'],
    'requires_prescription': ['requiere_receta', 'requiere receta', 'receta', 'requires_prescription'],
    'stock': ['stock', 'existencias', 'cantidad', 'inventario'],
    'description': ['descripcion', 'description'],
}

TEMPLATE_HEADERS = ['nombre', 'sku', 'codigo_barras', 'precio', 'categoria', 'requiere_receta', 'stock', 'descripcion']
TEMPLATE_EXAMPLE = [
    'Paracetamol 500mg (20 tabletas)', 'PAR-500', '7501234567890', 45.00,
    'Analgésicos y antiinflamatorios', 'no', 100, 'Analgésico y antipirético',
]

_TRUE_VALUES = {'si', 'sí', 'yes', 'true', '1', 'x'}


def _strip_accents(s: str) -> str:
    return (s.replace('á', 'a').replace('é', 'e').replace('í', 'i')
             .replace('ó', 'o').replace('ú', 'u').replace('ñ', 'n'))


def _normalize_header(h) -> str:
    return _strip_accents(str(h or '').strip().lower())


def _match_field(header) -> str:
    norm = _normalize_header(header)
    for field, aliases in COLUMN_ALIASES.items():
        if norm in (_normalize_header(a) for a in aliases):
            return field
    return None


def build_template_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Catálogo'
    ws.append(TEMPLATE_HEADERS)
    ws.append(TEMPLATE_EXAMPLE)
    for col in ws.columns:
        width = max(len(str(c.value)) for c in col if c.value is not None) + 2
        ws.column_dimensions[col[0].column_letter].width = max(width, 12)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_file(filename: str, content: bytes) -> list:
    """Returns a list of dicts with our normalized field names as keys, one per data row."""
    name = (filename or '').lower()
    rows = []
    if name.endswith('.xlsx'):
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return []
        field_by_col = {i: _match_field(h) for i, h in enumerate(header)}
        for raw in it:
            if raw is None or all(v is None for v in raw):
                continue
            rows.append({field_by_col[i]: v for i, v in enumerate(raw) if field_by_col.get(i)})
    elif name.endswith('.csv'):
        text = content.decode('utf-8-sig')
        reader = csv.reader(io.StringIO(text))
        header = next(reader, None)
        if not header:
            return []
        field_by_col = {i: _match_field(h) for i, h in enumerate(header)}
        for raw in reader:
            if not raw or all(not (v or '').strip() for v in raw):
                continue
            rows.append({field_by_col[i]: v for i, v in enumerate(raw) if field_by_col.get(i)})
    else:
        raise ValueError('Formato no soportado. Sube un archivo .xlsx o .csv')
    return rows


def _parse_bool(val) -> bool:
    if val is None:
        return False
    return str(val).strip().lower() in _TRUE_VALUES


def _parse_price(val):
    if val is None or str(val).strip() == '':
        return None
    if isinstance(val, (int, float)):
        return Decimal(str(val))
    cleaned = str(val).strip().replace('$', '').replace(',', '')
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def _parse_stock(val):
    if val is None or str(val).strip() == '':
        return None
    try:
        return int(float(str(val).strip()))
    except ValueError:
        return None


async def process_rows(db, rows: list, commit: bool) -> dict:
    """Matches each row to an existing product (by SKU, then barcode, then exact name)
    or flags it as a new product, computes a human-readable diff of what would change,
    and — only when commit=True — actually applies it. Each row is isolated in its own
    savepoint so one bad row (e.g. a duplicate barcode) doesn't abort the whole batch."""
    results = []
    categories_created = []

    cat_rows = (await db.execute(select(models.Category))).scalars().all()
    categories = {c.name.strip().lower(): c for c in cat_rows}

    products = (await db.execute(select(models.Product))).scalars().all()
    by_sku = {p.sku.strip().lower(): p for p in products if p.sku}
    by_barcode = {p.barcode.strip().lower(): p for p in products if p.barcode}
    by_name = {p.name.strip().lower(): p for p in products}

    inv_rows = (await db.execute(select(models.Inventory))).scalars().all()
    inventories = {inv.product_id: inv for inv in inv_rows}

    for idx, row in enumerate(rows, start=2):  # row 1 is the header
        name = (row.get('name') or '').strip() if row.get('name') else ''
        if not name:
            results.append({'row': idx, 'action': 'error', 'error': 'Falta el nombre del producto'})
            continue
        price = _parse_price(row.get('price'))
        if price is None:
            results.append({'row': idx, 'action': 'error', 'name': name, 'error': 'Precio inválido o faltante'})
            continue

        sku = (str(row.get('sku')).strip() if row.get('sku') else None) or None
        barcode = (str(row.get('barcode')).strip() if row.get('barcode') else None) or None
        category_name = (str(row.get('category')).strip() if row.get('category') else None) or None
        requires_prescription = _parse_bool(row.get('requires_prescription'))
        stock = _parse_stock(row.get('stock'))
        description = (str(row.get('description')).strip() if row.get('description') else None) or None

        existing = None
        if sku and sku.lower() in by_sku:
            existing = by_sku[sku.lower()]
        elif barcode and barcode.lower() in by_barcode:
            existing = by_barcode[barcode.lower()]
        elif name.lower() in by_name:
            existing = by_name[name.lower()]

        category_id = None
        category_note = None
        if category_name:
            key = category_name.lower()
            if key in categories:
                category_id = categories[key].id
            elif commit:
                slug = re.sub(r'[^a-z0-9]+', '-', _strip_accents(key)).strip('-') or None
                new_cat = models.Category(name=category_name, slug=slug)
                db.add(new_cat)
                await db.flush()
                categories[key] = new_cat
                category_id = new_cat.id
                categories_created.append(category_name)
            else:
                category_note = f'se creará la categoría "{category_name}"'

        row_result = {
            'row': idx, 'name': name, 'sku': sku, 'barcode': barcode,
            'price': float(price), 'category': category_name, 'category_note': category_note,
            'requires_prescription': requires_prescription, 'stock': stock,
        }

        try:
            if existing:
                inv = inventories.get(existing.id)
                changes = {}
                if float(existing.price) != float(price):
                    changes['precio'] = f'${float(existing.price):.2f} → ${float(price):.2f}'
                if sku and (existing.sku or None) != sku:
                    changes['sku'] = f'{existing.sku or "—"} → {sku}'
                if barcode and (existing.barcode or None) != barcode:
                    changes['codigo_barras'] = f'{existing.barcode or "—"} → {barcode}'
                if existing.requires_prescription != requires_prescription:
                    changes['receta'] = f'{"sí" if existing.requires_prescription else "no"} → {"sí" if requires_prescription else "no"}'
                if category_id and existing.category_id != category_id:
                    changes['categoria'] = f'→ {category_name}'
                if stock is not None and inv and inv.quantity != stock:
                    if stock < inv.reserved:
                        row_result['action'] = 'error'
                        row_result['error'] = f'El stock ({stock}) es menor a lo ya reservado en pedidos activos ({inv.reserved})'
                        results.append(row_result)
                        continue
                    changes['stock'] = f'{inv.quantity} → {stock}'

                row_result['action'] = 'update' if changes else 'sin_cambios'
                row_result['changes'] = changes

                if commit and changes:
                    async with db.begin_nested():
                        existing.price = price
                        if sku:
                            existing.sku = sku
                        if barcode:
                            existing.barcode = barcode
                        existing.requires_prescription = requires_prescription
                        if category_id:
                            existing.category_id = category_id
                        if description:
                            existing.description = description
                        if stock is not None and inv:
                            inv.quantity = stock
                        await db.flush()
            else:
                row_result['action'] = 'create'
                if commit:
                    async with db.begin_nested():
                        new_product = models.Product(
                            name=name, sku=sku, barcode=barcode, price=price,
                            requires_prescription=requires_prescription,
                            category_id=category_id, description=description,
                        )
                        db.add(new_product)
                        await db.flush()
                        db.add(models.Inventory(product_id=new_product.id, quantity=stock or 0, reserved=0))
                        await db.flush()
                    if sku:
                        by_sku[sku.lower()] = new_product
                    if barcode:
                        by_barcode[barcode.lower()] = new_product
                    by_name[name.lower()] = new_product
            results.append(row_result)
        except IntegrityError:
            row_result['action'] = 'error'
            row_result['error'] = 'El SKU o código de barras ya pertenece a otro producto'
            results.append(row_result)

    if commit:
        await db.commit()

    return {'rows': results, 'categories_created': categories_created}
